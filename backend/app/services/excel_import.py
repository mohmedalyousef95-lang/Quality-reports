import openpyxl

from ..database import SessionLocal
from ..models import School, ChecklistItem, NoteSuggestion
from ..constants import DEFAULT_CHECKLIST_ITEMS, DEFAULT_NOTE_SUGGESTIONS
from ..config import SEED_XLSX_PATH

COLUMNS = [
    "ministry_number",
    "name",
    "region",
    "address",
    "zone",
    "engineer",
    "supervisor",
    "lat",
    "lng",
    "building_type",
    "national_address",
    "ownership_type",
]


def import_schools(xlsx_path=SEED_XLSX_PATH, force: bool = False) -> int:
    db = SessionLocal()
    try:
        if not force and db.query(School).first() is not None:
            # Already populated — skip the 561-row upsert to keep boots fast.
            return db.query(School).count()
    finally:
        db.close()

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    db = SessionLocal()
    count = 0
    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = ws.iter_rows(min_row=2, max_col=len(COLUMNS), values_only=True)
            for row in rows:
                ministry_number = row[0]
                if not ministry_number:
                    continue
                ministry_number = str(ministry_number).strip()
                values = dict(zip(COLUMNS, row))
                values["ministry_number"] = ministry_number

                school = db.get(School, ministry_number)
                if school is None:
                    school = School(ministry_number=ministry_number)
                    db.add(school)
                for key, value in values.items():
                    setattr(school, key, value)
                count += 1
        db.commit()
    finally:
        db.close()
    return count


def seed_checklist_items() -> None:
    db = SessionLocal()
    try:
        existing = db.query(ChecklistItem).count()
        if existing > 0:
            return
        for category, items in DEFAULT_CHECKLIST_ITEMS.items():
            for position, label in enumerate(items):
                db.add(
                    ChecklistItem(category=category, label=label, position=position)
                )
        db.commit()
    finally:
        db.close()


def seed_note_suggestions() -> None:
    db = SessionLocal()
    try:
        existing = db.query(NoteSuggestion).count()
        if existing > 0:
            return
        for (category, item), texts in DEFAULT_NOTE_SUGGESTIONS.items():
            for position, text in enumerate(texts):
                db.add(
                    NoteSuggestion(
                        category=category, item=item, text=text, position=position
                    )
                )
        db.commit()
    finally:
        db.close()


if __name__ == "__main__":
    from ..database import Base, engine

    Base.metadata.create_all(bind=engine)
    n = import_schools()
    seed_checklist_items()
    print(f"Imported/updated {n} schools.")
