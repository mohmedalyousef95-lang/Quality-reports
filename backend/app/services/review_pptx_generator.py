from collections import defaultdict
from copy import deepcopy

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from ..config import TEMPLATE_PATH
from ..constants import (
    PHOTO_CATEGORIES,
    NOTE_CATEGORIES,
    PHOTO_CATEGORY_LABELS,
    NOTE_CATEGORY_LABELS,
    PHOTO_SLIDE_INDEX,
    NOTE_SLIDE_INDEX,
    REVIEW_RESOLVED_STATUS,
    REVIEW_RESPONSE_OPTIONS,
)

# This generator deliberately reuses pptx_generator's private helpers
# instead of duplicating them: those functions carry hard-won fixes for
# real PowerPoint-won't-open bugs this session (table element order,
# duplicate a16 tracking ids, zip partname collisions) — reimplementing
# any of them risks reintroducing the exact same defects from scratch.
from .pptx_generator import (
    _fill_cover,
    _force_title_font,
    _style_cell,
    _extract_locality,
    _prefetch_photo_bytes,
    _fill_photo_slide,
    _expand_notes_table_to_three_cols,
    _add_combined_notes_slides,
    _delete_slide,
    _move_slide,
    _add_fade_transitions,
    _renumber_slide_parts,
    _validate_deck_integrity,
)


def generate_review_pptx(school, review, output_path) -> None:
    prs = Presentation(TEMPLATE_PATH)
    original_report = review.report

    _fill_cover(prs, school, review)

    photo_layout = prs.slides[PHOTO_SLIDE_INDEX["ac"]].slide_layout

    src_slide = prs.slides[min(NOTE_SLIDE_INDEX.values())]
    notes_layout = src_slide.slide_layout
    src_table_el = next(sh for sh in src_slide.shapes if sh.has_table)._element
    table_template = deepcopy(src_table_el)
    _expand_notes_table_to_three_cols(
        table_template,
        title_label="متابعة الملاحظات",
        col_labels=("الملاحظة الأصلية", "الإجراء المتخذ / الملاحظة الحالية", "مدى التجاوب"),
    )

    # The template's fixed photo/notes slides are all replaced by slides
    # built dynamically below (before/after photo pairs, a notes table
    # keyed off review data) — delete them first so python-pptx's
    # slide-partname allocator keeps handing out gap-free numbers (see
    # _renumber_slide_parts's docstring for why order matters here).
    for idx in sorted(list(PHOTO_SLIDE_INDEX.values()) + list(NOTE_SLIDE_INDEX.values()), reverse=True):
        _delete_slide(prs, idx)

    before_photos_by_cat = defaultdict(list)
    for photo in sorted(original_report.photos, key=lambda p: p.position):
        before_photos_by_cat[photo.category].append(photo)

    after_photos_by_cat = defaultdict(list)
    for photo in sorted(review.photos, key=lambda p: p.position):
        after_photos_by_cat[photo.category].append(photo)

    photo_bytes = _prefetch_photo_bytes(list(original_report.photos) + list(review.photos))

    for category in PHOTO_CATEGORIES:
        before = before_photos_by_cat.get(category, [])
        after = after_photos_by_cat.get(category, [])
        if not before and not after:
            continue
        label = PHOTO_CATEGORY_LABELS[category]

        before_slide = prs.slides.add_slide(photo_layout)
        if before_slide.shapes.title is not None:
            before_slide.shapes.title.text = f"{label} — قبل"
            _force_title_font(before_slide.shapes.title)
        _fill_photo_slide(before_slide, before, photo_bytes)

        after_slide = prs.slides.add_slide(photo_layout)
        if after_slide.shapes.title is not None:
            after_slide.shapes.title.text = f"{label} — بعد"
            _force_title_font(after_slide.shapes.title)
        _fill_photo_slide(after_slide, after, photo_bytes)

    review_notes_by_category = defaultdict(list)
    for note in sorted(review.notes, key=lambda n: n.position):
        review_notes_by_category[note.category].append(note)

    _add_combined_notes_slides(
        prs,
        review_notes_by_category,
        notes_layout,
        table_template,
        slide_title="متابعة الملاحظات",
        status_of=lambda n: n.response_status or "لم تُحدَّد بعد",
    )

    # Order is now: cover, closing, photo pairs…, notes…
    # Move the closing slide back to the end, then insert the info/summary
    # slides right after the cover (summary first so the final read order
    # is: cover, info, summary, closing).
    _move_slide(prs, from_index=1, to_index=len(prs.slides) - 1)
    _add_review_summary_slide(prs, review)
    _add_review_info_slide(prs, school, review)

    _add_fade_transitions(prs)
    _renumber_slide_parts(prs)
    _validate_deck_integrity(prs)

    prs.save(output_path)


def _add_review_info_slide(prs, school, review) -> None:
    original_report = review.report

    region = school.region or ""
    locality = _extract_locality(school.address)
    location = " - ".join([x for x in [region, locality] if x]) or region

    rows = [
        ("اسم المدرسة", school.name or ""),
        ("الرقم الوزاري", school.ministry_number or ""),
        ("الموقع", location),
        ("الزون", school.zone or ""),
        ("مهندس الزون", school.engineer or ""),
        ("اسم المشرف", school.supervisor or ""),
        ("تاريخ الزيارة الأصلية", original_report.visit_date.strftime("%Y-%m-%d") if original_report.visit_date else ""),
        ("زائر الزيارة الأصلية", original_report.visitor_name or ""),
        ("تاريخ المراجعة الحالية", review.visit_date.strftime("%Y-%m-%d") if review.visit_date else ""),
        ("زائر المراجعة", review.visitor_name or ""),
    ]

    layout = prs.slides[-1].slide_layout  # "Title Only" (theme background)
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        title_shape = slide.shapes.title
        title_shape.text = "معلومات المراجعة"
        _force_title_font(title_shape)
        title_shape.left = 839788
        title_shape.top = 365125
        title_shape.width = 10515600
        title_shape.height = 722053

    n = len(rows)
    row_h, size = Inches(0.48), 13
    tbl_w = Inches(11.5)
    right_margin = Inches(0.92)
    top = Inches(1.35)
    tbl_h = row_h * n

    left = int(prs.slide_width - right_margin - tbl_w)
    graphic = slide.shapes.add_table(n, 2, left, int(top), int(tbl_w), int(tbl_h))
    table = graphic.table
    table.first_row = False
    table.horz_banding = False
    table.columns[0].width = int(tbl_w * 0.70)
    table.columns[1].width = int(tbl_w * 0.30)

    accent = RGBColor(0x00, 0x99, 0xA1)
    light = RGBColor(0xEC, 0xF3, 0xF2)
    white = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (label, value) in enumerate(rows):
        _style_cell(table.cell(i, 1), label, accent, light, bold=True, size=size, anchor=PP_ALIGN.CENTER)
        _style_cell(table.cell(i, 0), value, accent, white, bold=False, size=size, anchor=PP_ALIGN.CENTER)
        table.rows[i].height = int(row_h)

    _move_slide(prs, from_index=len(prs.slides) - 1, to_index=1)


def _add_review_summary_slide(prs, review) -> None:
    """A standalone slide summarising how many of this review's notes
    landed in each response-status bucket, plus the overall closure rate
    (resolved / total)."""
    counts = {opt: 0 for opt in REVIEW_RESPONSE_OPTIONS}
    unset = 0
    for note in review.notes:
        status = note.response_status or ""
        if status in counts:
            counts[status] += 1
        else:
            unset += 1
    total = len(review.notes)
    resolved = counts.get(REVIEW_RESOLVED_STATUS, 0)
    rate = round(resolved / total * 100) if total else 0

    rows = [("إجمالي الملاحظات", str(total))]
    for opt in REVIEW_RESPONSE_OPTIONS:
        rows.append((opt, str(counts[opt])))
    if unset:
        rows.append(("لم تُحدَّد بعد", str(unset)))

    layout = prs.slides[-1].slide_layout
    slide = prs.slides.add_slide(layout)
    if slide.shapes.title is not None:
        title_shape = slide.shapes.title
        title_shape.text = "ملخص التجاوب"
        _force_title_font(title_shape)
        title_shape.left = 839788
        title_shape.top = 365125
        title_shape.width = 10515600
        title_shape.height = 722053

    n = len(rows)
    row_h, size = Inches(0.60), 15
    tbl_w = Inches(11.5)
    right_margin = Inches(0.92)
    top = Inches(1.6)
    tbl_h = row_h * (n + 1)  # + the closure-rate banner row

    left = int(prs.slide_width - right_margin - tbl_w)
    graphic = slide.shapes.add_table(n + 1, 2, left, int(top), int(tbl_w), int(tbl_h))
    table = graphic.table
    table.first_row = False
    table.horz_banding = False
    table.columns[0].width = int(tbl_w * 0.30)
    table.columns[1].width = int(tbl_w * 0.70)

    accent = RGBColor(0x00, 0x99, 0xA1)
    light = RGBColor(0xEC, 0xF3, 0xF2)
    white = RGBColor(0xFF, 0xFF, 0xFF)

    for i, (label, value) in enumerate(rows):
        _style_cell(table.cell(i, 0), value, accent, light, bold=True, size=size, anchor=PP_ALIGN.CENTER)
        _style_cell(table.cell(i, 1), label, accent, white, bold=False, size=size, anchor=PP_ALIGN.CENTER)
        table.rows[i].height = int(row_h)

    rate_cell = table.cell(n, 0)
    rate_cell.merge(table.cell(n, 1))
    _style_cell(
        rate_cell,
        f"نسبة الإغلاق الإجمالية: {rate}٪",
        white,
        accent,
        bold=True,
        size=size + 4,
        anchor=PP_ALIGN.CENTER,
    )
    table.rows[n].height = int(row_h * 1.3)

    _move_slide(prs, from_index=len(prs.slides) - 1, to_index=1)
